"""Shared writer for the manuscript Supplementary Tables S2-S12 (one house style for every workbook).

    from supp_xlsx import write_table, OUT
    write_table(4, "atlas_ciliated_vs_secretory_markers",
                title="Supplementary Table 4. ...",
                readme=["Source: ...", "Method: ...", "Columns: ..."],
                sheets=[("All_DEGs", df1), ("Canonical_markers", df2)])

Writes OUT/SuppTable{n}_{slug}.xlsx (README sheet first, then data sheets: bold grey header, frozen header row,
auto column widths) and one CSV per data sheet: OUT/SuppTable{n}_{slug}_{sheet}.csv.
Table S1 uses its own writer (01_tableS1_atlas_studies.py).
"""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.environ.get("OUTPUT_ROOT", "output"), "final_panels/supplementary/tables")
HDR, FILL = Font(bold=True), PatternFill("solid", fgColor="EDEDED")


def _style(ws, nhdr=1, maxw=45):
    for r in range(1, nhdr + 1):
        for c in ws[r]:
            c.font = HDR; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = ws.cell(row=nhdr + 1, column=1)
    for i, col in enumerate(ws.columns, 1):
        w = max((len(str(c.value)) for c in col if c.value is not None), default=6)
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 8), maxw)


def write_table(n, slug, title, readme, sheets, csv=True):
    os.makedirs(OUT, exist_ok=True)
    base = f"{OUT}/SuppTable{n}_{slug}"
    wb = Workbook(); ws = wb.active; ws.title = "README"
    ws.append([title]); ws["A1"].font = Font(bold=True, size=12)
    for line in readme: ws.append([line])
    ws.column_dimensions["A"].width = 160
    for name, df in sheets:
        s = wb.create_sheet(name[:31]); s.append([str(c) for c in df.columns])
        for row in df.itertuples(index=False):
            s.append([None if (isinstance(v, float) and v != v) else v for v in row])
        _style(s)
        if csv: df.to_csv(f"{base}_{name}.csv", index=False)
    wb.save(f"{base}.xlsx")
    print(f"wrote {base}.xlsx  sheets: " + ", ".join(f"{nm} ({len(d)} rows)" for nm, d in sheets))
    return base

#!/usr/bin/env python3
"""
Propagate the metastatic_site correction (vazquez ascites: metastasis -> ascites)
into the ATLAS DOCUMENTATION artifacts, to match the corrected master atlas
(see fix_metastatic_site_vazquez.py). Documentation only; no analysis re-run.

Targets (each backed up first):
  1. pre-processing and integration/atlas_metadata.xlsx   (534 sample rows)
     rows where study==vazquez_garcia_2022 & anatomic_site==ascites &
     metastatic_site==metastasis  ->  metastatic_site = ascites
  2. output/supplemental/T1_atlas_metadata.csv            (per-cell export)
     same rule at the cell level.

Rule: anatomic_site is authoritative; vazquez ascites samples are the ascites
compartment. Only these cells/rows change.
"""
import os
import shutil
import pandas as pd
import openpyxl

ROOT = "/Volumes/CookLab/Sarah/202601_final_atlas"
XLSX = os.path.join(ROOT, "pre-processing and integration/atlas_metadata.xlsx")
T1 = os.path.join(ROOT, "output/supplemental/T1_atlas_metadata.csv")

STUDY, ANAT, FROM, TO = "vazquez_garcia_2022", "ascites", "metastasis", "ascites"


def fix_xlsx():
    bak = XLSX.replace(".xlsx", "_PRECORRECTION_backup.xlsx")
    if not os.path.exists(bak):
        shutil.copy2(XLSX, bak)
        print(f"  backup -> {os.path.basename(bak)}")
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    hdr = {c.value: c.column for c in ws[1]}
    cs, ca, cm = hdr["study"], hdr["anatomic_site"], hdr["metastatic_site"]
    n = 0
    for r in range(2, ws.max_row + 1):
        if (ws.cell(r, cs).value == STUDY and ws.cell(r, ca).value == ANAT
                and ws.cell(r, cm).value == FROM):
            ws.cell(r, cm).value = TO
            n += 1
    wb.save(XLSX)
    print(f"  atlas_metadata.xlsx: {n} sample rows recoded {FROM}->{TO}")


def fix_t1():
    if not os.path.exists(T1):
        print("  T1 not found, skipping")
        return
    bak = T1.replace(".csv", "_PRECORRECTION_backup.csv")
    if not os.path.exists(bak):
        shutil.copy2(T1, bak)
        print(f"  backup -> {os.path.basename(bak)}")
    d = pd.read_csv(T1, low_memory=False)
    mask = ((d["study"] == STUDY) & (d["anatomic_site"] == ANAT)
            & (d["metastatic_site"] == FROM))
    n = int(mask.sum())
    d.loc[mask, "metastatic_site"] = TO
    d.to_csv(T1, index=False)
    print(f"  T1_atlas_metadata.csv: {n} cells recoded {FROM}->{TO}")


if __name__ == "__main__":
    print("Correcting atlas documentation artifacts:")
    fix_xlsx()
    fix_t1()
    print("done.")
